"""
Comprehensive Unit Tests for the Mustaqil Ish Word (.docx) Document Generator.
Covers:
- Standard OTM margins (3.0cm left, 1.5cm right, 2.0cm top/bottom)
- Typography (Times New Roman 14pt, 1.5 line spacing, 1.25cm indent, Justified)
- Title page validation and error raising on missing metadata
- Native Word TOC field code presence
- Section styling (Headings 1 & 2)
- Zero placeholder text assertions
- Re-opening generated .docx files with python-docx to assert XML structure
"""

import io
import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docgen.mustaqil_ish import (
    MustaqilIshDocxBuilder,
    MustaqilIshDocument,
    TitlePageInfo,
    SectionItem,
    DEFAULT_PAGE_CONFIG,
)


class TestMustaqilIshDocxBuilder(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.builder = MustaqilIshDocxBuilder()

    def _create_sample_doc(self, title_info: TitlePageInfo, sections_count: int = 3) -> MustaqilIshDocument:
        sections = [
            SectionItem(
                title="KIRISH",
                paragraphs=[
                    "Ushbu mustaqil ish O'zbekiston iqtisodiyotining dolzarb masalalariga bag'ishlangan.",
                    "Tadqiqot maqsadi raqamlashtirish jarayonlarini chuqur tahlil qilishdan iborat."
                ],
                is_heading1=True,
                page_break_before=False
            )
        ]
        for i in range(1, sections_count + 1):
            sections.append(
                SectionItem(
                    title=f"{i}-BOB. IQTISODIY ISLOHOTLAR VA RIVOJLANISH BOSQICHLARI",
                    paragraphs=[
                        f"Ushbu {i}-bobda iqtisodiy o'sish dinamikasi va statistik ko'rsatkichlar tahlil qilinadi.",
                        "Bozor mexanizmlarining rivojlanishi yangi investitsiya imkoniyatlarini yaratmoqda."
                    ],
                    is_heading1=True,
                    page_break_before=True
                )
            )
            sections.append(
                SectionItem(
                    title=f"{i}.1-§. Tarmoqlararo integratsiya va samaradorlik",
                    paragraphs=[
                        "Samaradorlik ko'rsatkichlari korxonalarning raqobatbardoshligini belgilaydi."
                    ],
                    is_heading1=False,
                    page_break_before=False
                )
            )

        sections.append(
            SectionItem(
                title="XULOSA",
                paragraphs=[
                    "Olingan natijalar O'zbekistonda iqtisodiy modernizatsiya jadallashganini ko'rsatadi."
                ],
                is_heading1=True,
                page_break_before=True
            )
        )

        refs = [
            "O'zbekiston Respublikasi Prezidentining 2023-yil 11-sentabrdagi PF-158-son Farmoni.",
            "Mirziyoyev Sh.M. Yangi O'zbekiston taraqqiyot strategiyasi. -- Toshkent: O'zbekiston, 2022. -- B. 120-145.",
            "Statistika agentligi rasmiy portali. -- www.stat.uz."
        ]

        return MustaqilIshDocument(title_info=title_info, sections=sections, references=refs)

    def test_missing_title_info_raises_error(self):
        # Missing subject
        with self.assertRaises(ValueError):
            info = TitlePageInfo(
                university="",
                faculty="Iqtisodiyot",
                department="Moliya",
                subject="",
                topic="Raqamli iqtisodiyot",
                student_name="Azizbek",
                group_name="DI-101"
            )
            info.validate()

    def test_missing_student_raises_error(self):
        with self.assertRaises(ValueError):
            info = TitlePageInfo(
                university="TDIU",
                faculty="Iqtisodiyot",
                department="Moliya",
                subject="Makroiqtisodiyot",
                topic="Raqamli iqtisodiyot",
                student_name="",
                group_name="DI-101"
            )
            info.validate()

    def test_valid_title_info_passes(self):
        info = TitlePageInfo(
            university="Toshkent davlat iqtisodiyot universiteti",
            faculty="Iqtisodiyot va boshqaruv",
            department="Makroiqtisodiyot",
            subject="Makroiqtisodiyot",
            topic="Raqamli iqtisodiyotning rivojlanish istiqbollari",
            student_name="Aliyev Vali",
            group_name="MM-202",
            teacher_name="prof. Karimov A."
        )
        # Should not raise
        info.validate()

    def test_generate_default_document_bytes(self):
        info = TitlePageInfo(
            university="Toshkent davlat iqtisodiyot universiteti",
            faculty="Iqtisodiyot",
            department="Moliya",
            subject="Iqtisodiyot nazariyasi",
            topic="Bozor iqtisodiyoti asoslari",
            student_name="Sodiqov Jasur",
            group_name="IN-301",
            teacher_name="dots. Umarov B."
        )
        doc_data = self._create_sample_doc(info, sections_count=2)
        docx_bytes = self.builder.generate_docx(doc_data)

        self.assertIsInstance(docx_bytes, bytes)
        self.assertGreater(len(docx_bytes), 5000)

    def test_margins_exact_measurements(self):
        info = TitlePageInfo(
            university="O'zMU",
            faculty="Iqtisodiyot",
            department="Iqtisodiyot nazariyasi",
            subject="Iqtisodiyot",
            topic="Investitsiyalar",
            student_name="Nazarov D.",
            group_name="IQ-101"
        )
        doc_data = self._create_sample_doc(info, sections_count=1)
        docx_bytes = self.builder.generate_docx(doc_data)

        # Re-open with docx to assert XML geometry
        doc = docx.Document(io.BytesIO(docx_bytes))
        sec = doc.sections[0]

        self.assertAlmostEqual(sec.left_margin.cm, 3.0, places=2)
        self.assertAlmostEqual(sec.right_margin.cm, 1.5, places=2)
        self.assertAlmostEqual(sec.top_margin.cm, 2.0, places=2)
        self.assertAlmostEqual(sec.bottom_margin.cm, 2.0, places=2)
        self.assertTrue(sec.different_first_page_header_footer)

    def test_overridden_margins_config(self):
        custom_config = {
            "margin_left": Cm(2.5),
            "margin_right": Cm(1.0),
            "margin_top": Cm(2.0),
            "margin_bottom": Cm(2.0),
        }
        custom_builder = MustaqilIshDocxBuilder(custom_config=custom_config)
        info = TitlePageInfo(
            university="SamDU",
            faculty="Moliya",
            department="Bank ishi",
            subject="Bank ishi",
            topic="Kredit tizimi",
            student_name="Rustamov E.",
            group_name="BI-204"
        )
        doc_data = self._create_sample_doc(info, sections_count=1)
        docx_bytes = custom_builder.generate_docx(doc_data)

        doc = docx.Document(io.BytesIO(docx_bytes))
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.left_margin.cm, 2.5, places=2)
        self.assertAlmostEqual(sec.right_margin.cm, 1.0, places=2)

    def test_body_typography_specifications(self):
        info = TitlePageInfo(
            university="TATU",
            faculty="Dasturiy injiniring",
            department="Axborot texnologiyalari",
            subject="Sun'iy intellekt",
            topic="Mashinaviy o'rganish",
            student_name="Azizov B.",
            group_name="DI-401"
        )
        doc_data = self._create_sample_doc(info, sections_count=1)
        docx_bytes = self.builder.generate_docx(doc_data)

        doc = docx.Document(io.BytesIO(docx_bytes))
        # Find a body paragraph with justified text
        body_found = False
        for p in doc.paragraphs:
            if p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
                body_found = True
                self.assertEqual(p.paragraph_format.line_spacing, 1.5)
                self.assertAlmostEqual(p.paragraph_format.first_line_indent.cm, 1.25, places=2)
                for run in p.runs:
                    self.assertEqual(run.font.name, "Times New Roman")
                    self.assertEqual(run.font.size.pt, 14)
                break
        self.assertTrue(body_found, "Justified body paragraph not found")

    def test_toc_field_xml_present(self):
        info = TitlePageInfo(
            university="TDIU",
            faculty="Moliya",
            department="Moliya",
            subject="Moliya",
            topic="Davlat byudjeti",
            student_name="Karimov O.",
            group_name="MB-101"
        )
        doc_data = self._create_sample_doc(info, sections_count=2)
        docx_bytes = self.builder.generate_docx(doc_data)

        doc = docx.Document(io.BytesIO(docx_bytes))
        xml_str = doc._element.xml
        # Must contain TOC field
        self.assertIn("TOC", xml_str)

    def test_no_placeholders_in_document(self):
        info = TitlePageInfo(
            university="O'zbekiston Milliy Universiteti",
            faculty="Iqtisodiyot",
            department="Menejment",
            subject="Menejment",
            topic="Strategik boshqaruv",
            student_name="Yoqubov S.",
            group_name="SB-201",
            teacher_name="prof. Mahmudov T."
        )
        doc_data = self._create_sample_doc(info, sections_count=3)
        docx_bytes = self.builder.generate_docx(doc_data)

        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = " ".join([p.text for p in doc.paragraphs])

        self.assertNotIn("___________", full_text)
        self.assertNotIn("Lorem ipsum", full_text)
        self.assertNotIn("TODO", full_text)
        self.assertNotIn("FIXME", full_text)

    def test_large_document_8_sections(self):
        info = TitlePageInfo(
            university="JIDU",
            faculty="Xalqaro iqtisodiyot",
            department="Tashqi savdo",
            subject="Xalqaro savdo",
            topic="JST va O'zbekiston integratsiyasi",
            student_name="Sobirov A.",
            group_name="XS-301"
        )
        doc_data = self._create_sample_doc(info, sections_count=8)
        docx_bytes = self.builder.generate_docx(doc_data)

        self.assertGreater(len(docx_bytes), 15000)
        doc = docx.Document(io.BytesIO(docx_bytes))
        self.assertGreater(len(doc.paragraphs), 30)


    def test_default_title_page_omits_university_header(self):
        info = TitlePageInfo(
            university="Toshkent davlat iqtisodiyot universiteti",
            faculty="Iqtisodiyot",
            department="Kafedra",
            subject="Mikroiqtisodiyot",
            topic="Korxona xarajatlari",
            student_name="Aziz Raxmatullayev",
            group_name="IQ-101"
        )
        doc_data = self._create_sample_doc(info, sections_count=1)
        docx_bytes = self.builder.generate_docx(doc_data)
        doc = docx.Document(io.BytesIO(docx_bytes))
        first_page_paras = [p.text for p in doc.paragraphs[:15] if p.text.strip()]
        first_page_text = " ".join(first_page_paras)
        self.assertNotIn("VAZIRLIGI", first_page_text)
        self.assertNotIn("TOSHKENT DAVLAT", first_page_text)
        self.assertIn("MUSTAQIL ISH", first_page_text)
        self.assertIn("Mikroiqtisodiyot", first_page_text)


if __name__ == "__main__":
    unittest.main()
